"""
core/loaders.py

Extracts clean text from uploaded PDF, DOCX, or TXT files.
Every function raises a `DocumentLoadError` with a human-readable message on
failure — the UI layer catches this and shows it directly, so error text
here should be end-user friendly, not a stack trace.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field


class DocumentLoadError(Exception):
    """Raised when a document cannot be parsed into usable text."""


@dataclass
class LoadedDocument:
    filename: str
    text: str
    page_count: int = 0
    warnings: list[str] = field(default_factory=list)


def _clean_text(raw: str) -> str:
    """Normalize whitespace without destroying paragraph/clause breaks."""
    if not raw:
        return ""
    # Collapse excessive blank lines but keep paragraph structure
    lines = [line.rstrip() for line in raw.splitlines()]
    cleaned_lines = []
    blank_run = 0
    for line in lines:
        if line.strip() == "":
            blank_run += 1
            if blank_run <= 1:
                cleaned_lines.append("")
        else:
            blank_run = 0
            cleaned_lines.append(line)
    return "\n".join(cleaned_lines).strip()


def load_pdf(file_bytes: bytes, filename: str) -> LoadedDocument:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise DocumentLoadError(
            "PDF support isn't installed. Run: pip install pypdf"
        ) from e

    try:
        reader = PdfReader(io.BytesIO(file_bytes))
    except Exception as e:
        raise DocumentLoadError(
            f"'{filename}' could not be opened — it may be corrupted, "
            f"password-protected, or not a valid PDF. ({e.__class__.__name__})"
        ) from e

    if reader.is_encrypted:
        try:
            reader.decrypt("")  # try empty password first
        except Exception:
            raise DocumentLoadError(
                f"'{filename}' is password-protected. Please upload an "
                "unlocked version of the document."
            )

    warnings = []
    text_parts = []
    for i, page in enumerate(reader.pages):
        try:
            page_text = page.extract_text() or ""
        except Exception:
            page_text = ""
            warnings.append(f"Page {i + 1} could not be read and was skipped.")
        text_parts.append(page_text)

    full_text = _clean_text("\n\n".join(text_parts))

    if not full_text or len(full_text) < 20:
        raise DocumentLoadError(
            f"'{filename}' appears to contain no extractable text. It may be "
            "a scanned image PDF (ClauseShield doesn't perform OCR in this "
            "build) or an empty file."
        )

    return LoadedDocument(
        filename=filename,
        text=full_text,
        page_count=len(reader.pages),
        warnings=warnings,
    )


def load_docx(file_bytes: bytes, filename: str) -> LoadedDocument:
    try:
        import docx
    except ImportError as e:
        raise DocumentLoadError(
            "DOCX support isn't installed. Run: pip install python-docx"
        ) from e

    try:
        document = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise DocumentLoadError(
            f"'{filename}' could not be opened — it may be corrupted or not "
            f"a valid Word document. ({e.__class__.__name__})"
        ) from e

    paragraphs = [p.text for p in document.paragraphs]

    # Also pull text out of tables, since contracts often use them for
    # payment schedules / deliverables
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                paragraphs.append(row_text)

    full_text = _clean_text("\n".join(paragraphs))

    if not full_text or len(full_text) < 20:
        raise DocumentLoadError(
            f"'{filename}' appears to be empty or contains no readable text."
        )

    return LoadedDocument(filename=filename, text=full_text, page_count=0)


def load_txt(file_bytes: bytes, filename: str) -> LoadedDocument:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            raw = file_bytes.decode(encoding)
            break
        except (UnicodeDecodeError, LookupError):
            continue
    else:
        raise DocumentLoadError(
            f"'{filename}' could not be decoded as text. Please save it as "
            "UTF-8 plain text and try again."
        )

    full_text = _clean_text(raw)

    if not full_text or len(full_text) < 20:
        raise DocumentLoadError(f"'{filename}' is empty.")

    return LoadedDocument(filename=filename, text=full_text, page_count=0)


def load_document(file_bytes: bytes, filename: str) -> LoadedDocument:
    """Dispatches to the correct loader based on file extension."""
    if not file_bytes:
        raise DocumentLoadError(f"'{filename}' is empty (0 bytes).")

    lower = filename.lower()
    if lower.endswith(".pdf"):
        return load_pdf(file_bytes, filename)
    if lower.endswith(".docx"):
        return load_docx(file_bytes, filename)
    if lower.endswith(".txt"):
        return load_txt(file_bytes, filename)

    raise DocumentLoadError(
        f"Unsupported file type for '{filename}'. Please upload a PDF, "
        "DOCX, or TXT file."
    )
